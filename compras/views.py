from decimal import Decimal, InvalidOperation
import os
import base64

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.conf import settings
from django.db import transaction
from django.db.models import Q, Sum, Avg
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template import loader
from django.template.loader import render_to_string
from django.urls import reverse_lazy
from django.utils.decorators import method_decorator
from django.views import generic
from django.utils import timezone

from datetime import date, datetime
import io

# ── Excel ──
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── PDF ──
from xhtml2pdf import pisa

from flor.models import Flor
from producto.models import Producto
from proveedores.models import Proveedor
from core.notifications import crear_notificacion, crear_notificacion_stock
from usuarios.decorators import panel_login_required

from .forms import CompraForm
from .models import Compra, DetalleCompra


# ────────────────────────────────────────────────
#  HELPERS INTERNOS
# ────────────────────────────────────────────────

def _parse_item_id(raw_item_id):
    tipo_raw, raw_id = (raw_item_id or "").split("-", 1)
    item_id = int(raw_id)
    if tipo_raw == "F":
        return "FLOR", item_id
    if tipo_raw == "P":
        return "PRODUCTO", item_id
    raise ValueError("Tipo de item invalido")


def _parse_detalles_compra(request):
    item_ids   = request.POST.getlist("item_id[]")
    precios    = request.POST.getlist("precio[]")
    cantidades = request.POST.getlist("cantidad[]")

    detalles = []
    for idx, (item_id_raw, precio_raw, cantidad_raw) in enumerate(zip(item_ids, precios, cantidades), start=1):
        item_id_raw  = str(item_id_raw).strip()
        precio_raw   = str(precio_raw).strip()
        cantidad_raw = str(cantidad_raw).strip()

        if not item_id_raw and not precio_raw and not cantidad_raw:
            continue
        if not item_id_raw:
            raise ValueError(f"Articulo {idx}: Debe seleccionar un item.")
        if not precio_raw or not cantidad_raw:
            raise ValueError(f"Articulo {idx}: Debe completar precio y cantidad.")

        try:
            tipo_item, item_pk = _parse_item_id(item_id_raw)
        except (ValueError, TypeError):
            raise ValueError(f"Articulo {idx}: Item invalido.")

        try:
            precio_normalizado = precio_raw
            if "," in precio_raw:
                precio_normalizado = precio_raw.replace(".", "").replace(",", ".")
            precio_limpio   = Decimal(precio_normalizado)
            cantidad_limpia = int(cantidad_raw)
        except (InvalidOperation, ValueError):
            raise ValueError(f"Articulo {idx}: Formato invalido en precio o cantidad.")

        if precio_limpio <= 0:
            raise ValueError(f"Articulo {idx}: El precio debe ser mayor a 0.")
        if cantidad_limpia <= 0:
            raise ValueError(f"Articulo {idx}: La cantidad debe ser mayor a 0.")

        detalles.append({
            "tipo_item": tipo_item,
            "item_pk":   item_pk,
            "cantidad":  cantidad_limpia,
            "precio":    precio_limpio,
        })

    if not detalles:
        raise ValueError("Debe agregar al menos un articulo con precio y cantidad validos.")
    return detalles


def _obtener_items_posteados_compra(request):
    item_ids = request.POST.getlist("item_id[]")
    precios = request.POST.getlist("precio[]")
    cantidades = request.POST.getlist("cantidad[]")

    total_filas = max(len(item_ids), len(precios), len(cantidades))
    items = []

    for idx in range(total_filas):
        item_id_raw = (item_ids[idx] if idx < len(item_ids) else "").strip()
        precio_raw = (precios[idx] if idx < len(precios) else "").strip()
        cantidad_raw = (cantidades[idx] if idx < len(cantidades) else "").strip()

        if not item_id_raw and not precio_raw and not cantidad_raw:
            continue

        tipo_item = ""
        item_pk = ""
        if "-" in item_id_raw:
            prefijo, pk_raw = item_id_raw.split("-", 1)
            if prefijo == "F":
                tipo_item = "FLOR"
            elif prefijo == "P":
                tipo_item = "PRODUCTO"
            item_pk = pk_raw.strip()

        items.append(
            {
                "item_id": item_id_raw,
                "tipo_item": tipo_item,
                "item_pk": item_pk,
                "precio": precio_raw,
                "cantidad": cantidad_raw,
            }
        )

    return items


def _bloquear_item(tipo_item, item_pk):
    if tipo_item == "FLOR":
        return Flor.objects.select_for_update().get(pk=item_pk)
    return Producto.objects.select_for_update().get(pk=item_pk)


def _sumar_stock_item(tipo_item, item_pk, cantidad):
    item = _bloquear_item(tipo_item, item_pk)
    item.cantidad += cantidad
    item.save(update_fields=["cantidad"])
    crear_notificacion_stock(item.nombre, item.cantidad, "Compra")


def _restar_stock_item(tipo_item, item_pk, cantidad, contexto):
    item = _bloquear_item(tipo_item, item_pk)
    if item.cantidad < cantidad:
        raise ValueError(
            f"No hay stock suficiente para ajustar {item.nombre} en {contexto}. "
            f"Disponible: {item.cantidad}, requerido: {cantidad}."
        )
    item.cantidad -= cantidad
    item.save(update_fields=["cantidad"])
    crear_notificacion_stock(item.nombre, item.cantidad, contexto)


def _thin_border():
    side = Side(style="thin", color="D1D5DB")
    return Border(left=side, right=side, top=side, bottom=side)


def _get_logo_base64():
    """Carga el logo desde compras/static/img/LogoAE.png como base64."""
    logo_path = os.path.join(settings.BASE_DIR, "compras", "static", "img", "LogoAE.png")
    if os.path.exists(logo_path):
        with open(logo_path, "rb") as f:
            return base64.b64encode(f.read()).decode("utf-8")
    return ""


# ────────────────────────────────────────────────
#  LISTA DE COMPRAS
# ────────────────────────────────────────────────

@login_required
@panel_login_required
def compras_list(request):
    lista_compras = Compra.objects.select_related("proveedor", "usuario").prefetch_related(
        "detalles__flor", "detalles__producto"
    )

    q                = request.GET.get("q", "").strip()
    proveedor_nombre = request.GET.get("proveedor_nombre", "").strip()
    fecha_desde      = request.GET.get("fecha_desde", "")
    precio_min       = request.GET.get("precio_min", "").strip()
    precio_max       = request.GET.get("precio_max", "").strip()

    def _parse_decimal(raw):
        valor = (raw or "").strip()
        if not valor:
            return None
        try:
            if "," in valor:
                valor = valor.replace(".", "").replace(",", ".")
            return Decimal(valor)
        except (InvalidOperation, ValueError, TypeError):
            return None

    if q:
        filtros_q = (
            Q(descripcion__icontains=q) |
            Q(proveedor__nombre_proveedor__icontains=q) |
            Q(proveedor__numero_documento__icontains=q)
        )
        if q.isdigit():
            filtros_q |= Q(id=int(q))
        lista_compras = lista_compras.filter(filtros_q)

    if proveedor_nombre:
        lista_compras = lista_compras.filter(
            proveedor__nombre_proveedor__icontains=proveedor_nombre
        )

    if fecha_desde:
        try:
            fecha = datetime.strptime(fecha_desde, "%Y-%m-%d").date()
            lista_compras = lista_compras.filter(fecha_emision__gte=fecha)
        except ValueError:
            fecha_desde = ""

    precio_min_val = _parse_decimal(precio_min)
    precio_max_val = _parse_decimal(precio_max)
    if precio_min and precio_min_val is None:
        precio_min = ""
    if precio_max and precio_max_val is None:
        precio_max = ""
    if precio_min_val is not None:
        lista_compras = lista_compras.filter(total_compra__gte=precio_min_val)
    if precio_max_val is not None:
        lista_compras = lista_compras.filter(total_compra__lte=precio_max_val)

    proveedores       = Proveedor.objects.all().order_by("nombre_proveedor")
    total_compras     = lista_compras.count()
    monto_total       = lista_compras.aggregate(Sum("total_compra"))["total_compra__sum"] or 0
    total_proveedores = lista_compras.values("proveedor").distinct().count()

    hoy         = date.today()
    inicio_mes  = hoy.replace(day=1)
    compras_mes = lista_compras.filter(
        fecha_emision__gte=inicio_mes, fecha_emision__lte=hoy
    ).count()

    anio_actual       = timezone.now().year
    anios_raw         = Compra.objects.dates("fecha_emision", "year", order="DESC")
    anios_disponibles = [d.year for d in anios_raw] or [anio_actual]

    template = loader.get_template("lista_compra.html")
    context = {
        "compras":                 lista_compras,
        "query":                   q,
        "total_compras":           total_compras,
        "monto_total":             monto_total,
        "total_proveedores":       total_proveedores,
        "compras_mes":             compras_mes,
        "proveedores":             proveedores,
        "proveedor_nombre_filtro": proveedor_nombre,
        "fecha_desde_filtro":      fecha_desde,
        "precio_min_filtro":       precio_min,
        "precio_max_filtro":       precio_max,
        "resultados_filtrados":    lista_compras.count(),
        "hay_filtros":             any([q, proveedor_nombre, fecha_desde, precio_min, precio_max]),
        "anios_disponibles":       anios_disponibles,
        "anio_actual":             anio_actual,
    }
    return HttpResponse(template.render(context, request))


# ────────────────────────────────────────────────
#  DETALLE DE COMPRA
# ────────────────────────────────────────────────

@login_required
@panel_login_required
def compra_detail(request, id):
    una_compra = get_object_or_404(
        Compra.objects.select_related("proveedor", "usuario").prefetch_related(
            "detalles__flor", "detalles__producto"
        ),
        id=id,
    )
    una_compra.calcular_totales()
    template = loader.get_template("compra_detail.html")
    return HttpResponse(
        template.render(
            {"compra": una_compra, "detalles": una_compra.detalles.all()}, request
        )
    )


# ────────────────────────────────────────────────
#  REPORTE DE COMPRAS
# ────────────────────────────────────────────────

MESES = {
    1: "Enero",   2: "Febrero",  3: "Marzo",      4: "Abril",
    5: "Mayo",    6: "Junio",    7: "Julio",       8: "Agosto",
    9: "Septiembre", 10: "Octubre", 11: "Noviembre", 12: "Diciembre",
}


def reporte(request):
    anio    = request.GET.get("anio")
    mes     = request.GET.get("mes")
    formato = request.GET.get("formato", "html")

    try:
        anio = int(anio)
    except (TypeError, ValueError):
        anio = timezone.now().year

    qs = (
        Compra.objects
        .filter(fecha_emision__year=anio)
        .select_related("proveedor")
        .prefetch_related("detalles")
    )

    mes_nombre = None
    if mes:
        try:
            mes_int    = int(mes)
            mes_nombre = MESES.get(mes_int, "")
            qs         = qs.filter(fecha_emision__month=mes_int)
        except ValueError:
            mes = None

    qs = qs.order_by("fecha_emision")

    agg         = qs.aggregate(total=Sum("total_compra"), promedio=Avg("total_compra"))
    monto_total = agg["total"]    or Decimal("0")
    promedio    = agg["promedio"] or Decimal("0")
    total_reg   = qs.count()
    total_prov  = qs.values("proveedor").distinct().count()

    periodo   = f"{mes_nombre} {anio}".strip() if mes_nombre else str(anio)
    fecha_gen = datetime.now().strftime("%d/%m/%Y %H:%M")
    usuario   = (
        request.user.get_full_name() or request.user.username
        if request.user.is_authenticated else ""
    )

    context = {
        "compras":           qs,
        "periodo":           periodo,
        "mes_nombre":        mes_nombre,
        "mes":               mes or "",
        "anio":              anio,
        "monto_total":       monto_total,
        "promedio_compra":   promedio,
        "total_registros":   total_reg,
        "total_proveedores": total_prov,
        "fecha_generacion":  fecha_gen,
        "usuario":           usuario,
        "hay_filtros":       bool(mes_nombre),
        "proveedor_filtro":  request.GET.get("proveedor_nombre", ""),
        "logo_base64":       _get_logo_base64(),
    }

    if formato == "excel":
        return _reporte_excel(qs, periodo, monto_total)

    if formato == "pdf":
        html_string = render_to_string("reporte.html", context, request=request)
        buffer      = io.BytesIO()
        pisa_status = pisa.CreatePDF(html_string, dest=buffer)
        if pisa_status.err:
            return HttpResponse("Error al generar el PDF", status=500)
        buffer.seek(0)
        filename = f"reporte_compras_{periodo.replace(' ', '_')}.pdf"
        response = HttpResponse(buffer, content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    return render(request, "reporte.html", context)


# ────────────────────────────────────────────────
#  EXCEL HELPER
# ────────────────────────────────────────────────

def _reporte_excel(qs, periodo, total_general):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Compras"

    PURPLE = "7C3AED"
    LIGHT  = "F3F0FF"

    ws.merge_cells("A1:G1")
    ws["A1"] = f"Reporte de Compras — Artes & Estilos | Período: {periodo}"
    ws["A1"].font      = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill      = PatternFill("solid", fgColor=PURPLE)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = ["#", "Fecha", "Proveedor", "Documento", "Descripción", "Total", "Ítems"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font      = Font(bold=True, color="FFFFFF", size=10)
        cell.fill      = PatternFill("solid", fgColor=PURPLE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = _thin_border()

    for i, c in enumerate(qs, start=3):
        fill = (PatternFill("solid", fgColor=LIGHT) if i % 2 == 0
                else PatternFill("solid", fgColor="FFFFFF"))
        row_data = [
            c.id,
            c.fecha_emision.strftime("%d/%m/%Y"),
            c.proveedor.nombre_proveedor if c.proveedor else "N/A",
            c.proveedor.numero_documento if c.proveedor else "—",
            c.descripcion or "Sin descripción",
            float(c.total_compra),
            c.detalles.count(),
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.fill      = fill
            cell.border    = _thin_border()
            cell.alignment = Alignment(vertical="center")
            if col == 6:
                cell.number_format = '"$"#,##0.00'
                cell.alignment     = Alignment(horizontal="right", vertical="center")

    last = ws.max_row + 1
    ws.cell(row=last, column=5, value="TOTAL GENERAL").font = Font(bold=True)
    tot = ws.cell(row=last, column=6, value=float(total_general))
    tot.font          = Font(bold=True)
    tot.number_format = '"$"#,##0.00'
    tot.alignment     = Alignment(horizontal="right")
    tot.fill          = PatternFill("solid", fgColor="EDE9FE")
    for col in range(1, 8):
        ws.cell(row=last, column=col).border = _thin_border()

    # ✅ Anchos de columna usando letras (evita error MergedCell)
    for col, w in zip(['A', 'B', 'C', 'D', 'E', 'F', 'G'], [6, 12, 28, 18, 45, 14, 7]):
        ws.column_dimensions[col].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"reporte_compras_{periodo.replace(' ', '_')}.xlsx"
    response = HttpResponse(
        buffer,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ────────────────────────────────────────────────
#  CLASS-BASED VIEWS
# ────────────────────────────────────────────────

@method_decorator(panel_login_required, name="dispatch")
class CompraCreateView(LoginRequiredMixin, generic.CreateView):
    model         = Compra
    form_class    = CompraForm
    template_name = "crear_compra.html"
    success_url   = reverse_lazy("compras:lista_compra")
    login_url     = "usuarios:login"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["proveedores"] = Proveedor.objects.filter(activo=True).order_by("nombre_proveedor")
        context["flores"]      = Flor.objects.all().order_by("nombre")
        context["productos"]   = Producto.objects.all().order_by("nombre")
        return context

    def form_valid(self, form):
        try:
            detalles = _parse_detalles_compra(self.request)
        except ValueError as exc:
            messages.error(self.request, str(exc))
            return self.form_invalid(form)
        try:
            with transaction.atomic():
                compra              = form.save(commit=False)
                compra.subtotal     = 0
                compra.total_compra = 0
                compra.usuario      = self.request.user
                compra.save()
                for data in detalles:
                    detalle = DetalleCompra(
                        compra=compra, tipo_item=data["tipo_item"],
                        cantidad=data["cantidad"], precio=data["precio"],
                    )
                    if data["tipo_item"] == "FLOR":
                        detalle.flor = Flor.objects.get(pk=data["item_pk"])
                    else:
                        detalle.producto = Producto.objects.get(pk=data["item_pk"])
                    detalle.save()
                    _sumar_stock_item(data["tipo_item"], data["item_pk"], data["cantidad"])
                compra.calcular_totales()
                crear_notificacion(
                    categoria="movimiento",
                    estilo="success",
                    titulo="Compra creada",
                    mensaje=f"Se registro la compra #{compra.id} con {len(detalles)} item(s).",
                )
            messages.success(
                self.request,
                f"Compra registrada exitosamente con {len(detalles)} item(s).",
            )
            return redirect(self.success_url)
        except (Flor.DoesNotExist, Producto.DoesNotExist):
            messages.error(self.request, "Uno de los items seleccionados ya no existe.")
            return self.form_invalid(form)
        except ValueError as exc:
            messages.error(self.request, str(exc))
            return self.form_invalid(form)
        except Exception as exc:
            messages.error(self.request, f"Error al guardar la compra: {exc}")
            return self.form_invalid(form)

    def form_invalid(self, form):
        if form.errors:
            messages.error(self.request, "Por favor, corrija los errores en el formulario.")
        if self.request.method == "POST":
            context = self.get_context_data(
                form=form,
                posted_items=_obtener_items_posteados_compra(self.request),
            )
            return self.render_to_response(context)
        return super().form_invalid(form)


@method_decorator(panel_login_required, name="dispatch")
class CompraUpdateView(LoginRequiredMixin, generic.UpdateView):
    model         = Compra
    form_class    = CompraForm
    template_name = "editar_compra.html"
    success_url   = reverse_lazy("compras:lista_compra")
    pk_url_kwarg  = "compra_id"
    login_url     = "usuarios:login"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["proveedores"] = Proveedor.objects.filter(activo=True).order_by("nombre_proveedor")
        context["detalles"]    = self.object.detalles.select_related("flor", "producto")
        context["flores"]      = Flor.objects.all().order_by("nombre")
        context["productos"]   = Producto.objects.all().order_by("nombre")
        return context

    def form_valid(self, form):
        try:
            nuevos_detalles = _parse_detalles_compra(self.request)
        except ValueError as exc:
            messages.error(self.request, str(exc))
            return self.form_invalid(form)
        compra = self.get_object()
        try:
            with transaction.atomic():
                compra = form.save()

                detalles_actuales = list(compra.detalles.select_related("flor", "producto"))

                stock_actual_por_item = {}
                for detalle in detalles_actuales:
                    item_pk = detalle.flor_id if detalle.tipo_item == "FLOR" else detalle.producto_id
                    if not item_pk:
                        continue
                    key = (detalle.tipo_item, item_pk)
                    stock_actual_por_item[key] = stock_actual_por_item.get(key, 0) + int(detalle.cantidad)

                stock_nuevo_por_item = {}
                for data in nuevos_detalles:
                    key = (data["tipo_item"], data["item_pk"])
                    stock_nuevo_por_item[key] = stock_nuevo_por_item.get(key, 0) + int(data["cantidad"])

                for key, cantidad_actual in stock_actual_por_item.items():
                    cantidad_nueva = stock_nuevo_por_item.get(key, 0)
                    if cantidad_actual > cantidad_nueva:
                        tipo_item, item_pk = key
                        _restar_stock_item(
                            tipo_item, item_pk,
                            cantidad_actual - cantidad_nueva,
                            "la edicion de compra",
                        )

                for key, cantidad_nueva in stock_nuevo_por_item.items():
                    cantidad_actual = stock_actual_por_item.get(key, 0)
                    if cantidad_nueva > cantidad_actual:
                        tipo_item, item_pk = key
                        _sumar_stock_item(tipo_item, item_pk, cantidad_nueva - cantidad_actual)

                compra.detalles.all().delete()
                for data in nuevos_detalles:
                    detalle = DetalleCompra(
                        compra=compra, tipo_item=data["tipo_item"],
                        cantidad=data["cantidad"], precio=data["precio"],
                    )
                    if data["tipo_item"] == "FLOR":
                        detalle.flor = Flor.objects.get(pk=data["item_pk"])
                    else:
                        detalle.producto = Producto.objects.get(pk=data["item_pk"])
                    detalle.save()
                    _sumar_stock_item(data["tipo_item"], data["item_pk"], data["cantidad"])

                compra.calcular_totales()
                crear_notificacion(
                    categoria="movimiento",
                    estilo="info",
                    titulo="Compra actualizada",
                    mensaje=f"Se actualizo la compra #{compra.id} con {len(nuevos_detalles)} item(s).",
                )
            messages.success(
                self.request,
                f"Compra actualizada exitosamente con {len(nuevos_detalles)} item(s).",
            )
            return redirect(self.success_url)
        except (Flor.DoesNotExist, Producto.DoesNotExist):
            messages.error(self.request, "Uno de los items seleccionados ya no existe.")
            return self.form_invalid(form)
        except ValueError as exc:
            messages.error(self.request, str(exc))
            return self.form_invalid(form)
        except Exception as exc:
            messages.error(self.request, f"Error al actualizar la compra: {exc}")
            return self.form_invalid(form)

    def form_invalid(self, form):
        if form.errors:
            messages.error(self.request, "Por favor, corrija los errores en el formulario.")
        if self.request.method == "POST":
            context = self.get_context_data(
                form=form,
                posted_items=_obtener_items_posteados_compra(self.request),
            )
            return self.render_to_response(context)
        return super().form_invalid(form)


@method_decorator(panel_login_required, name="dispatch")
class CompraDeleteView(LoginRequiredMixin, generic.DeleteView):
    model         = Compra
    template_name = "eliminar_compra.html"
    success_url   = reverse_lazy("compras:lista_compra")
    pk_url_kwarg  = "compra_id"
    login_url     = "usuarios:login"

    def post(self, request, *args, **kwargs):
        compra = self.get_object()
        try:
            with transaction.atomic():
                for detalle in list(compra.detalles.select_related("flor", "producto")):
                    item_pk = (
                        detalle.flor_id if detalle.tipo_item == "FLOR"
                        else detalle.producto_id
                    )
                    if item_pk:
                        _restar_stock_item(
                            detalle.tipo_item, item_pk,
                            detalle.cantidad, "la eliminacion de compra",
                        )
                compra.delete()
                crear_notificacion(
                    categoria="movimiento",
                    estilo="error",
                    titulo="Compra eliminada",
                    mensaje=f"Se elimino la compra #{compra.id}.",
                )
            messages.success(request, f"La compra {compra.id} ha sido eliminada exitosamente.")
        except ValueError as exc:
            messages.error(request, str(exc))
        except Exception as exc:
            messages.error(request, f"No se pudo eliminar la compra: {exc}")
        return redirect(self.success_url)