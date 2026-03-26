import base64
import io
import os
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.conf import settings
from django.core.files.base import ContentFile
from django.db.models import Q, Sum
from django.http import HttpResponse
from django.shortcuts import render
from django.urls import reverse_lazy
from django.utils import timezone
from django.views import generic

import openpyxl # type: ignore
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side # type: ignore

from core.notifications import crear_notificacion
from compras.models import Compra
from .forms import FlorForm
from .models import Flor

TIPOS_DISPLAY = dict(Flor._meta.get_field("tipo_flor").choices) if hasattr(Flor, '_meta') else {}


# ────────────────────────────────────────────────
#  HELPERS
# ────────────────────────────────────────────────

def _parse_decimal(valor):
    if valor is None:
        return None
    limpio = str(valor).strip()
    if not limpio:
        return None
    normalizado = limpio.replace(".", "").replace(",", ".")
    try:
        return Decimal(normalizado)
    except (InvalidOperation, ValueError):
        return None


def _procesar_imagen(request, nombre_flor):
    cropped_data = request.POST.get("cropped_image_data", "").strip()
    if cropped_data:
        try:
            encabezado, imgstr = cropped_data.split(";base64,")
            ext = encabezado.split("/")[-1]
            nombre_archivo = f"flor_{nombre_flor}.{ext}"
            return ContentFile(base64.b64decode(imgstr), name=nombre_archivo)
        except Exception:
            pass
    return request.FILES.get("imagen") or None


# ✅ Después — agrega LogoAE.png
def _get_logo_base64():
    candidatos = []
    for static_dir in getattr(settings, "STATICFILES_DIRS", []):
        candidatos.append(os.path.join(static_dir, "img", "LogoAE.png"))
        candidatos.append(os.path.join(static_dir, "img", "logo.png"))
        candidatos.append(os.path.join(static_dir, "img", "logo.jpg"))
    static_root = getattr(settings, "STATIC_ROOT", None)
    if static_root:
        candidatos.append(os.path.join(static_root, "img", "LogoAE.png"))
        candidatos.append(os.path.join(static_root, "img", "logo.png"))
        candidatos.append(os.path.join(static_root, "img", "logo.jpg"))
    for path in candidatos:
        if os.path.exists(path):
            with open(path, "rb") as f:
                return base64.b64encode(f.read()).decode("utf-8")
    return None


def _thin_border():
    side = Side(style="thin", color="D1D5DB")
    return Border(left=side, right=side, top=side, bottom=side)


# ────────────────────────────────────────────────
#  LISTA DE FLORES
# ────────────────────────────────────────────────

class FlorListView(LoginRequiredMixin, UserPassesTestMixin, generic.ListView):
    model = Flor
    template_name = "flor/lista.html"
    context_object_name = "flores"
    login_url = "/panel/login/"

    def test_func(self):
        return self.request.user.is_staff

    def get_queryset(self):
        queryset = Flor.objects.all()

        self.q              = self.request.GET.get("q", "").strip()
        self.tipo           = self.request.GET.get("tipo", "").strip()
        self.nivel_stock    = self.request.GET.get("nivel_stock", "").strip()
        self.precio_min_raw = self.request.GET.get("precio_min", "").strip()
        self.precio_max_raw = self.request.GET.get("precio_max", "").strip()

        if self.q:
            queryset = queryset.filter(
                Q(nombre__icontains=self.q) |
                Q(descripcion__icontains=self.q) |
                Q(tipo_flor__icontains=self.q)
            )

        if self.tipo:
            queryset = queryset.filter(tipo_flor=self.tipo)

        if self.nivel_stock == "bajo":
            queryset = queryset.filter(cantidad__lte=10)
        elif self.nivel_stock == "medio":
            queryset = queryset.filter(cantidad__gte=11, cantidad__lte=30)
        elif self.nivel_stock == "alto":
            queryset = queryset.filter(cantidad__gt=30)

        precio_min = _parse_decimal(self.precio_min_raw)
        precio_max = _parse_decimal(self.precio_max_raw)

        if precio_min is not None and precio_max is not None and precio_min > precio_max:
            precio_min, precio_max = precio_max, precio_min
            self.precio_min_raw = str(precio_min)
            self.precio_max_raw = str(precio_max)
            messages.info(
                self.request,
                "Se ajusto el rango de precios porque el minimo era mayor al maximo.",
            )

        if precio_min is not None:
            queryset = queryset.filter(precio__gte=precio_min)
        if precio_max is not None:
            queryset = queryset.filter(precio__lte=precio_max)

        return queryset.order_by("-id")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        total_flores = Flor.objects.count()
        total_stock  = Flor.objects.aggregate(total=Sum('cantidad'))['total'] or 0
        bajo_stock   = Flor.objects.filter(cantidad__lte=10).count()
        tipos_usados = Flor.objects.values('tipo_flor').distinct().count()

        anio_actual       = timezone.now().year
        fechas            = Compra.objects.dates("fecha_emision", "year", order="DESC")
        anios_disponibles = [d.year for d in fechas]
        if not anios_disponibles:
            anios_disponibles = [anio_actual]
        elif anio_actual not in anios_disponibles:
            anios_disponibles.insert(0, anio_actual)

        context.update({
            'busqueda':             getattr(self, 'q', ''),
            'tipo_filtro':          getattr(self, 'tipo', ''),
            'nivel_stock_filtro':   getattr(self, 'nivel_stock', ''),
            'precio_min_filtro':    getattr(self, 'precio_min_raw', ''),
            'precio_max_filtro':    getattr(self, 'precio_max_raw', ''),
            'tipos_filtro':         Flor._meta.get_field('tipo_flor').choices,
            'tipos_reporte':        Flor._meta.get_field('tipo_flor').choices,
            'total_flores':         total_flores,
            'total_stock':          total_stock,
            'bajo_stock':           bajo_stock,
            'tipos_usados':         tipos_usados,
            'resultados_filtrados': context['flores'].count(),
            'hay_filtros': any([
                getattr(self, 'q', ''),
                getattr(self, 'tipo', ''),
                getattr(self, 'nivel_stock', ''),
                getattr(self, 'precio_min_raw', ''),
                getattr(self, 'precio_max_raw', ''),
            ]),
            'anios_disponibles': anios_disponibles,
            'anio_actual':       anio_actual,
        })
        return context


# ────────────────────────────────────────────────
#  REPORTE DE FLORES
# ────────────────────────────────────────────────

def reporte_flor(request):
    tipo_filtro = request.GET.get("tipo", "").strip()
    formato     = request.GET.get("formato", "html").strip()

    tipos_dict  = dict(Flor._meta.get_field("tipo_flor").choices)
    tipo_nombre = tipos_dict.get(tipo_filtro, "")

    qs = Flor.objects.all().order_by("tipo_flor", "nombre")
    if tipo_filtro:
        qs = qs.filter(tipo_flor=tipo_filtro)

    flores_lista = []
    for f in qs:
        f.valor_total = f.precio * f.cantidad
        flores_lista.append(f)

    total_flores     = qs.count()
    total_stock      = qs.aggregate(t=Sum("cantidad"))["t"] or 0
    bajo_stock       = qs.filter(cantidad__lte=10).count()
    valor_inventario = sum(f.valor_total for f in flores_lista)

    resumen_tipos = []
    if not tipo_filtro:
        for clave, nombre in Flor._meta.get_field("tipo_flor").choices:
            grupo = [f for f in flores_lista if f.tipo_flor == clave]
            if not grupo:
                continue
            resumen_tipos.append({
                "nombre":              nombre,
                "cantidad_variedades": len(grupo),
                "total_stock":         sum(f.cantidad for f in grupo),
                "precio_promedio":     sum(f.precio for f in grupo) / len(grupo),
                "valor_total":         sum(f.valor_total for f in grupo),
            })

    fecha_gen = timezone.now().strftime("%d/%m/%Y %H:%M")
    usuario   = (
        request.user.get_full_name() or request.user.username
        if request.user.is_authenticated else "—"
    )

    if formato == "excel":
        return _reporte_flor_excel(flores_lista, resumen_tipos, tipo_nombre, valor_inventario)

    context = {
        "flores":            flores_lista,
        "resumen_tipos":     resumen_tipos,
        "tipo_filtro":       tipo_filtro,
        "tipo_nombre":       tipo_nombre,
        "total_flores":      total_flores,
        "total_stock":       total_stock,
        "bajo_stock":        bajo_stock,
        "valor_inventario":  valor_inventario,
        "fecha_generacion":  fecha_gen,
        "usuario":           usuario,
        "logo_base64":       _get_logo_base64(),
    }
    return render(request, "flor/reporte_flor.html", context)


# ────────────────────────────────────────────────
#  EXCEL HELPER
# ────────────────────────────────────────────────

def _reporte_flor_excel(flores_lista, resumen_tipos, tipo_nombre, valor_inventario):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Flores"

    PINK  = "D4006E"
    LIGHT = "FDF0F7"

    ws.merge_cells("A1:G1")
    titulo = "Inventario de Flores — Artes & Estilos"
    if tipo_nombre:
        titulo += f" | Tipo: {tipo_nombre}"
    ws["A1"] = titulo
    ws["A1"].font      = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill      = PatternFill("solid", fgColor=PINK)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = ["#", "Nombre", "Tipo", "Descripción", "Precio", "Stock", "Valor total"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font      = Font(bold=True, color="FFFFFF", size=10)
        cell.fill      = PatternFill("solid", fgColor=PINK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = _thin_border()

    for i, f in enumerate(flores_lista, start=3):
        fill = (PatternFill("solid", fgColor=LIGHT) if i % 2 == 0
                else PatternFill("solid", fgColor="FFFFFF"))
        row_data = [
            i - 2,
            f.nombre,
            f.get_tipo_flor_display(),
            f.descripcion or "—",
            float(f.precio),
            f.cantidad,
            float(f.valor_total),
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.fill      = fill
            cell.border    = _thin_border()
            cell.alignment = Alignment(vertical="center")
            if col in (5, 7):
                cell.number_format = '"$"#,##0.00'
                cell.alignment     = Alignment(horizontal="right", vertical="center")

    last = ws.max_row + 1
    ws.cell(row=last, column=5, value="VALOR TOTAL INVENTARIO").font = Font(bold=True)
    tot = ws.cell(row=last, column=7, value=float(valor_inventario))
    tot.font          = Font(bold=True)
    tot.number_format = '"$"#,##0.00'
    tot.alignment     = Alignment(horizontal="right")
    tot.fill          = PatternFill("solid", fgColor="FFE5F4")
    for col in range(1, 8):
        ws.cell(row=last, column=col).border = _thin_border()

    # ✅ Sin .column_letter — usa letras directamente
    for col, w in zip(['A', 'B', 'C', 'D', 'E', 'F', 'G'], [5, 25, 15, 35, 12, 8, 14]):
        ws.column_dimensions[col].width = w

    # ── Hoja 2: Resumen por tipo ─────────────────
    if resumen_tipos:
        ws2 = wb.create_sheet("Por Tipo")
        ws2.merge_cells("A1:E1")
        ws2["A1"] = "Resumen por Tipo de Flor — Artes & Estilos"
        ws2["A1"].font      = Font(bold=True, size=13, color="FFFFFF")
        ws2["A1"].fill      = PatternFill("solid", fgColor=PINK)
        ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[1].height = 24

        headers2 = ["Tipo", "Variedades", "Unidades", "Precio prom.", "Valor total"]
        for col, h in enumerate(headers2, 1):
            cell = ws2.cell(row=2, column=col, value=h)
            cell.font      = Font(bold=True, color="FFFFFF", size=10)
            cell.fill      = PatternFill("solid", fgColor=PINK)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = _thin_border()

        for i, t in enumerate(resumen_tipos, start=3):
            fill = (PatternFill("solid", fgColor=LIGHT) if i % 2 == 0
                    else PatternFill("solid", fgColor="FFFFFF"))
            row_data2 = [
                t["nombre"], t["cantidad_variedades"],
                t["total_stock"], float(t["precio_promedio"]), float(t["valor_total"]),
            ]
            for col, val in enumerate(row_data2, 1):
                cell = ws2.cell(row=i, column=col, value=val)
                cell.fill      = fill
                cell.border    = _thin_border()
                if col in (4, 5):
                    cell.number_format = '"$"#,##0.00'
                    cell.alignment     = Alignment(horizontal="right", vertical="center")

        # ✅ Sin .column_letter — usa letras directamente
        for col, w in zip(['A', 'B', 'C', 'D', 'E'], [20, 12, 12, 14, 14]):
            ws2.column_dimensions[col].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"reporte_flores{'_' + tipo_nombre if tipo_nombre else ''}.xlsx"
    response = HttpResponse(
        buffer,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ────────────────────────────────────────────────
#  CRUD VIEWS
# ────────────────────────────────────────────────

class FlorCreateView(LoginRequiredMixin, UserPassesTestMixin, generic.CreateView):
    model         = Flor
    form_class    = FlorForm
    template_name = "flor/form.html"
    success_url   = reverse_lazy("flor:lista")
    login_url     = "/panel/login/"

    def test_func(self):
        return self.request.user.is_staff

    def form_valid(self, form):
        imagen = _procesar_imagen(self.request, form.cleaned_data.get("nombre") or "flor")
        if imagen:
            form.instance.imagen = imagen
        messages.success(self.request, "Flor creada correctamente.")
        crear_notificacion(
            categoria="movimiento", estilo="success",
            titulo="Flor creada",
            mensaje=f"Se creo la flor {form.instance.nombre}.",
        )
        return super().form_valid(form)


class FlorUpdateView(LoginRequiredMixin, UserPassesTestMixin, generic.UpdateView):
    model         = Flor
    form_class    = FlorForm
    template_name = "flor/form.html"
    success_url   = reverse_lazy("flor:lista")
    login_url     = "/panel/login/"

    def test_func(self):
        return self.request.user.is_staff

    def form_valid(self, form):
        imagen = _procesar_imagen(self.request, form.cleaned_data.get("nombre") or "flor")
        if imagen:
            form.instance.imagen = imagen
        messages.success(self.request, "Flor actualizada correctamente.")
        crear_notificacion(
            categoria="movimiento", estilo="info",
            titulo="Flor actualizada",
            mensaje=f"Se actualizo la flor {form.instance.nombre}.",
        )
        return super().form_valid(form)


class FlorDetailView(LoginRequiredMixin, UserPassesTestMixin, generic.DetailView):
    model               = Flor
    template_name       = "flor/detalle.html"
    context_object_name = "flor"
    login_url           = "/panel/login/"

    def test_func(self):
        return self.request.user.is_staff


class FlorDeleteView(LoginRequiredMixin, UserPassesTestMixin, generic.DeleteView):
    model         = Flor
    template_name = "flor/confirm_delete.html"
    success_url   = reverse_lazy("flor:lista")
    login_url     = "/panel/login/"

    def test_func(self):
        return self.request.user.is_staff

    def form_valid(self, form):
        nombre = self.object.nombre
        crear_notificacion(
            categoria="movimiento", estilo="error",
            titulo="Flor eliminada",
            mensaje=f"Se elimino la flor {nombre}.",
        )
        messages.success(self.request, "Flor eliminada correctamente.")
        return super().form_valid(form)