import base64
import io
import os
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.conf import settings
from django.core.files.base import ContentFile
from django.db.models import Q, Sum
from django.http import HttpResponse
from django.shortcuts import render
from django.urls import reverse_lazy
from django.utils import timezone
from django.views import generic

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from core.notifications import crear_notificacion
from .forms import ProductoForm
from .models import Producto


# ────────────────────────────────────────────────
#  HELPERS
# ────────────────────────────────────────────────

def _parse_decimal(valor):
    if valor is None:
        return None
    limpio = str(valor).strip()
    if not limpio:
        return None
    normalizado = limpio.replace(".", "").replace(",", ".")
    try:
        return Decimal(normalizado)
    except (InvalidOperation, ValueError):
        return None


def _procesar_imagen(request, nombre_producto):
    cropped_data = request.POST.get("cropped_image_data", "").strip()
    if cropped_data:
        try:
            encabezado, imgstr = cropped_data.split(";base64,")
            ext = encabezado.split("/")[-1]
            nombre_archivo = f"producto_{nombre_producto}.{ext}"
            return ContentFile(base64.b64decode(imgstr), name=nombre_archivo)
        except Exception:
            pass
    return request.FILES.get("imagen") or None


# ✅ Después — agrega LogoAE.png
def _get_logo_base64():
    candidatos = []
    for static_dir in getattr(settings, "STATICFILES_DIRS", []):
        candidatos.append(os.path.join(static_dir, "img", "LogoAE.png"))
        candidatos.append(os.path.join(static_dir, "img", "logo.png"))
        candidatos.append(os.path.join(static_dir, "img", "logo.jpg"))
    static_root = getattr(settings, "STATIC_ROOT", None)
    if static_root:
        candidatos.append(os.path.join(static_root, "img", "LogoAE.png"))
        candidatos.append(os.path.join(static_root, "img", "logo.png"))
        candidatos.append(os.path.join(static_root, "img", "logo.jpg"))
    for path in candidatos:
        if os.path.exists(path):
            with open(path, "rb") as f:
                return base64.b64encode(f.read()).decode("utf-8")
    return None


def _thin_border():
    side = Side(style="thin", color="D1D5DB")
    return Border(left=side, right=side, top=side, bottom=side)


# ────────────────────────────────────────────────
#  LISTA DE PRODUCTOS
# ────────────────────────────────────────────────

class ProductoListView(LoginRequiredMixin, UserPassesTestMixin, generic.ListView):
    model = Producto
    template_name = "producto/lista.html"
    context_object_name = "productos"
    login_url = "/panel/login/"

    def test_func(self):
        return self.request.user.is_staff

    def get_queryset(self):
        queryset = Producto.objects.all()

        self.q              = self.request.GET.get("q", "").strip()
        self.tipo           = self.request.GET.get("tipo", "").strip()
        self.nivel_stock    = self.request.GET.get("nivel_stock", "").strip()
        self.precio_min_raw = self.request.GET.get("precio_min", "").strip()
        self.precio_max_raw = self.request.GET.get("precio_max", "").strip()

        if self.q:
            queryset = queryset.filter(
                Q(nombre__icontains=self.q) |
                Q(descripcion__icontains=self.q) |
                Q(tipo_producto__icontains=self.q)
            )

        if self.tipo:
            queryset = queryset.filter(tipo_producto=self.tipo)

        if self.nivel_stock == "bajo":
            queryset = queryset.filter(cantidad__lte=10)
        elif self.nivel_stock == "medio":
            queryset = queryset.filter(cantidad__gte=11, cantidad__lte=30)
        elif self.nivel_stock == "alto":
            queryset = queryset.filter(cantidad__gt=30)

        precio_min = _parse_decimal(self.precio_min_raw)
        precio_max = _parse_decimal(self.precio_max_raw)

        if precio_min is not None and precio_max is not None and precio_min > precio_max:
            precio_min, precio_max = precio_max, precio_min
            self.precio_min_raw = str(precio_min)
            self.precio_max_raw = str(precio_max)
            messages.info(
                self.request,
                "Se ajusto el rango de precios porque el minimo era mayor al maximo.",
            )

        if precio_min is not None:
            queryset = queryset.filter(precio__gte=precio_min)
        if precio_max is not None:
            queryset = queryset.filter(precio__lte=precio_max)

        return queryset.order_by("-id")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        total_productos = Producto.objects.count()
        total_stock     = Producto.objects.aggregate(total=Sum("cantidad"))["total"] or 0
        bajo_stock      = Producto.objects.filter(cantidad__lte=10).count()
        tipos_usados    = Producto.objects.values("tipo_producto").distinct().count()

        context.update({
            "busqueda":             getattr(self, "q", ""),
            "tipo_filtro":          getattr(self, "tipo", ""),
            "nivel_stock_filtro":   getattr(self, "nivel_stock", ""),
            "precio_min_filtro":    getattr(self, "precio_min_raw", ""),
            "precio_max_filtro":    getattr(self, "precio_max_raw", ""),
            "tipos_filtro":         Producto._meta.get_field("tipo_producto").choices,
            "tipos_reporte":        Producto._meta.get_field("tipo_producto").choices,
            "total_productos":      total_productos,
            "total_stock":          total_stock,
            "bajo_stock":           bajo_stock,
            "tipos_usados":         tipos_usados,
            "resultados_filtrados": context["productos"].count(),
            "hay_filtros": any([
                getattr(self, "q", ""),
                getattr(self, "tipo", ""),
                getattr(self, "nivel_stock", ""),
                getattr(self, "precio_min_raw", ""),
                getattr(self, "precio_max_raw", ""),
            ]),
        })
        return context


# ────────────────────────────────────────────────
#  REPORTE DE PRODUCTOS
#  GET ?tipo=<valor>&formato=html|excel
# ────────────────────────────────────────────────

def reporte_producto(request):
    tipo_filtro = request.GET.get("tipo", "").strip()
    formato     = request.GET.get("formato", "html").strip()

    tipos_dict  = dict(Producto._meta.get_field("tipo_producto").choices)
    tipo_nombre = tipos_dict.get(tipo_filtro, "")

    qs = Producto.objects.all().order_by("tipo_producto", "nombre")
    if tipo_filtro:
        qs = qs.filter(tipo_producto=tipo_filtro)

    productos_lista = []
    for p in qs:
        p.valor_total = p.precio * p.cantidad
        productos_lista.append(p)

    total_productos  = qs.count()
    total_stock      = qs.aggregate(t=Sum("cantidad"))["t"] or 0
    bajo_stock       = qs.filter(cantidad__lte=10).count()
    valor_inventario = sum(p.valor_total for p in productos_lista)

    resumen_tipos = []
    if not tipo_filtro:
        for clave, nombre in Producto._meta.get_field("tipo_producto").choices:
            grupo = [p for p in productos_lista if p.tipo_producto == clave]
            if not grupo:
                continue
            resumen_tipos.append({
                "nombre":              nombre,
                "cantidad_variedades": len(grupo),
                "total_stock":         sum(p.cantidad for p in grupo),
                "precio_promedio":     sum(p.precio for p in grupo) / len(grupo),
                "valor_total":         sum(p.valor_total for p in grupo),
            })

    fecha_gen = timezone.now().strftime("%d/%m/%Y %H:%M")
    usuario   = (
        request.user.get_full_name() or request.user.username
        if request.user.is_authenticated else "—"
    )

    if formato == "excel":
        return _reporte_producto_excel(productos_lista, resumen_tipos, tipo_nombre, valor_inventario)

    context = {
        "productos":         productos_lista,
        "resumen_tipos":     resumen_tipos,
        "tipo_filtro":       tipo_filtro,
        "tipo_nombre":       tipo_nombre,
        "total_productos":   total_productos,
        "total_stock":       total_stock,
        "bajo_stock":        bajo_stock,
        "valor_inventario":  valor_inventario,
        "fecha_generacion":  fecha_gen,
        "usuario":           usuario,
        "logo_base64":       _get_logo_base64(),
    }
    return render(request, "producto/reporte_producto.html", context)


# ────────────────────────────────────────────────
#  EXCEL HELPER
# ────────────────────────────────────────────────

def _reporte_producto_excel(productos_lista, resumen_tipos, tipo_nombre, valor_inventario):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"

    TEAL  = "0D9488"
    LIGHT = "F0FDFA"

    ws.merge_cells("A1:G1")
    titulo = "Inventario de Productos — Artes & Estilos"
    if tipo_nombre:
        titulo += f" | Categoría: {tipo_nombre}"
    ws["A1"] = titulo
    ws["A1"].font      = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill      = PatternFill("solid", fgColor=TEAL)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = ["#", "Nombre", "Categoría", "Descripción", "Precio", "Stock", "Valor total"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font      = Font(bold=True, color="FFFFFF", size=10)
        cell.fill      = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = _thin_border()

    for i, p in enumerate(productos_lista, start=3):
        fill = (PatternFill("solid", fgColor=LIGHT) if i % 2 == 0
                else PatternFill("solid", fgColor="FFFFFF"))
        row_data = [
            i - 2,
            p.nombre,
            p.get_tipo_producto_display(),
            p.descripcion or "—",
            float(p.precio),
            p.cantidad,
            float(p.valor_total),
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.fill      = fill
            cell.border    = _thin_border()
            cell.alignment = Alignment(vertical="center")
            if col in (5, 7):
                cell.number_format = '"$"#,##0.00'
                cell.alignment     = Alignment(horizontal="right", vertical="center")

    last = ws.max_row + 1
    ws.cell(row=last, column=5, value="VALOR TOTAL INVENTARIO").font = Font(bold=True)
    tot = ws.cell(row=last, column=7, value=float(valor_inventario))
    tot.font          = Font(bold=True)
    tot.number_format = '"$"#,##0.00'
    tot.alignment     = Alignment(horizontal="right")
    tot.fill          = PatternFill("solid", fgColor="CCFBF1")
    for col in range(1, 8):
        ws.cell(row=last, column=col).border = _thin_border()

    # ✅ Letras directas — evita error MergedCell
    for col, w in zip(["A", "B", "C", "D", "E", "F", "G"], [5, 25, 18, 35, 12, 8, 14]):
        ws.column_dimensions[col].width = w

    # ── Hoja 2: Resumen por categoría ────────────
    if resumen_tipos:
        ws2 = wb.create_sheet("Por Categoría")
        ws2.merge_cells("A1:E1")
        ws2["A1"] = "Resumen por Categoría de Producto — Artes & Estilos"
        ws2["A1"].font      = Font(bold=True, size=13, color="FFFFFF")
        ws2["A1"].fill      = PatternFill("solid", fgColor=TEAL)
        ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[1].height = 24

        headers2 = ["Categoría", "Variedades", "Unidades", "Precio prom.", "Valor total"]
        for col, h in enumerate(headers2, 1):
            cell = ws2.cell(row=2, column=col, value=h)
            cell.font      = Font(bold=True, color="FFFFFF", size=10)
            cell.fill      = PatternFill("solid", fgColor=TEAL)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = _thin_border()

        for i, t in enumerate(resumen_tipos, start=3):
            fill = (PatternFill("solid", fgColor=LIGHT) if i % 2 == 0
                    else PatternFill("solid", fgColor="FFFFFF"))
            row_data2 = [
                t["nombre"], t["cantidad_variedades"],
                t["total_stock"], float(t["precio_promedio"]), float(t["valor_total"]),
            ]
            for col, val in enumerate(row_data2, 1):
                cell = ws2.cell(row=i, column=col, value=val)
                cell.fill      = fill
                cell.border    = _thin_border()
                if col in (4, 5):
                    cell.number_format = '"$"#,##0.00'
                    cell.alignment     = Alignment(horizontal="right", vertical="center")

        # ✅ Letras directas — evita error MergedCell
        for col, w in zip(["A", "B", "C", "D", "E"], [22, 12, 12, 14, 14]):
            ws2.column_dimensions[col].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"reporte_productos{'_' + tipo_nombre if tipo_nombre else ''}.xlsx"
    response = HttpResponse(
        buffer,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ────────────────────────────────────────────────
#  CRUD VIEWS (sin cambios)
# ────────────────────────────────────────────────

class ProductoCreateView(LoginRequiredMixin, UserPassesTestMixin, generic.CreateView):
    model         = Producto
    form_class    = ProductoForm
    template_name = "producto/form.html"
    success_url   = reverse_lazy("producto:lista")
    login_url     = "/panel/login/"

    def test_func(self):
        return self.request.user.is_staff

    def form_valid(self, form):
        imagen = _procesar_imagen(self.request, form.cleaned_data.get("nombre") or "producto")
        if imagen:
            form.instance.imagen = imagen
        messages.success(self.request, "Producto creado correctamente.")
        crear_notificacion(
            categoria="movimiento", estilo="success",
            titulo="Producto creado",
            mensaje=f"Se creo el producto {form.instance.nombre}.",
        )
        return super().form_valid(form)


class ProductoUpdateView(LoginRequiredMixin, UserPassesTestMixin, generic.UpdateView):
    model         = Producto
    form_class    = ProductoForm
    template_name = "producto/form.html"
    success_url   = reverse_lazy("producto:lista")
    login_url     = "/panel/login/"

    def test_func(self):
        return self.request.user.is_staff

    def form_valid(self, form):
        imagen = _procesar_imagen(self.request, form.cleaned_data.get("nombre") or "producto")
        if imagen:
            form.instance.imagen = imagen
        messages.success(self.request, "Producto actualizado correctamente.")
        crear_notificacion(
            categoria="movimiento", estilo="info",
            titulo="Producto actualizado",
            mensaje=f"Se actualizo el producto {form.instance.nombre}.",
        )
        return super().form_valid(form)


class ProductoDetailView(LoginRequiredMixin, UserPassesTestMixin, generic.DetailView):
    model               = Producto
    template_name       = "producto/detalle.html"
    context_object_name = "producto"
    login_url           = "/panel/login/"

    def test_func(self):
        return self.request.user.is_staff


class ProductoDeleteView(LoginRequiredMixin, UserPassesTestMixin, generic.DeleteView):
    model         = Producto
    template_name = "producto/confirm_delete.html"
    success_url   = reverse_lazy("producto:lista")
    login_url     = "/panel/login/"

    def test_func(self):
        return self.request.user.is_staff

    def form_valid(self, form):
        nombre = self.object.nombre
        crear_notificacion(
            categoria="movimiento", estilo="error",
            titulo="Producto eliminado",
            mensaje=f"Se elimino el producto {nombre}.",
        )
        messages.success(self.request, "Producto eliminado correctamente.")
        return super().form_valid(form)