import base64
import io
import os
from decimal import Decimal, InvalidOperation
from datetime import date, datetime

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Q, Sum, Avg, Count
from django.http import JsonResponse, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from clientes.models import Cliente
from flor.models import Flor
from producto.models import Producto
from core.notifications import crear_notificacion, crear_notificacion_stock
from usuarios.decorators import panel_login_required

from .forms import VentaForm
from .models import DetalleVenta, Venta


# ────────────────────────────────────────────────
#  CONSTANTES
# ────────────────────────────────────────────────

MESES = {
    1: "Enero",   2: "Febrero",  3: "Marzo",      4: "Abril",
    5: "Mayo",    6: "Junio",    7: "Julio",       8: "Agosto",
    9: "Septiembre", 10: "Octubre", 11: "Noviembre", 12: "Diciembre",
}

TIPO_VENTA_DISPLAY = {
    "BP": "Bajo Pedido",
    "EI": "Entrega Inmediata",
}


# ────────────────────────────────────────────────
#  HELPERS
# ────────────────────────────────────────────────

def _parse_item_id(raw_item_id):
    tipo_raw, raw_id = (raw_item_id or "").split("-", 1)
    item_id = int(raw_id)
    if tipo_raw == "F":
        return "FLOR", item_id
    if tipo_raw == "P":
        return "PRODUCTO", item_id
    raise ValueError("Tipo de item invalido")


def _parse_detalles_venta(request):
    arreglo_ids = request.POST.getlist("arreglo_id[]")
    cantidades  = request.POST.getlist("cantidad[]")
    precios     = request.POST.getlist("precio[]")

    detalles = []
    for idx, (arreglo_id, cantidad_raw, precio_raw) in enumerate(
        zip(arreglo_ids, cantidades, precios), start=1
    ):
        arreglo_id   = str(arreglo_id).strip()
        cantidad_raw = str(cantidad_raw).strip()
        precio_raw   = str(precio_raw).strip()

        if not arreglo_id and not cantidad_raw and not precio_raw:
            continue
        if not arreglo_id:
            raise ValueError(f"Item {idx}: Debes seleccionar un flor o producto.")

        try:
            tipo_item, item_pk = _parse_item_id(arreglo_id)
        except (ValueError, TypeError):
            raise ValueError(f"Item {idx}: Identificador invalido.")

        try:
            cantidad = int(cantidad_raw)
        except ValueError:
            raise ValueError(f"Item {idx}: Formato invalido en cantidad.")

        precio = None
        if precio_raw:
            try:
                precio = Decimal(precio_raw)
            except InvalidOperation:
                raise ValueError(f"Item {idx}: Formato invalido en precio.")

        if not precio or precio <= 0:
            try:
                if tipo_item == "FLOR":
                    precio = Flor.objects.get(pk=item_pk).precio
                else:
                    precio = Producto.objects.get(pk=item_pk).precio
            except (Flor.DoesNotExist, Producto.DoesNotExist):
                raise ValueError(f"Item {idx}: El item seleccionado ya no existe.")

        if cantidad <= 0:
            raise ValueError(f"Item {idx}: La cantidad debe ser mayor a 0.")
        if precio <= 0:
            raise ValueError(f"Item {idx}: El precio debe ser mayor a 0.")

        detalles.append({
            "tipo_item": tipo_item,
            "item_pk":   item_pk,
            "cantidad":  cantidad,
            "precio":    precio,
        })

    if not detalles:
        raise ValueError("Debes agregar al menos un item a la venta.")
    return detalles


def _obtener_items_posteados_venta(request):
    arreglo_ids = request.POST.getlist("arreglo_id[]")
    cantidades = request.POST.getlist("cantidad[]")
    precios = request.POST.getlist("precio[]")

    total_filas = max(len(arreglo_ids), len(cantidades), len(precios))
    items = []

    for idx in range(total_filas):
        arreglo_id_raw = (arreglo_ids[idx] if idx < len(arreglo_ids) else "").strip()
        cantidad_raw = (cantidades[idx] if idx < len(cantidades) else "").strip()
        precio_raw = (precios[idx] if idx < len(precios) else "").strip()

        if not arreglo_id_raw and not cantidad_raw and not precio_raw:
            continue

        tipo_item = ""
        item_pk = ""
        if "-" in arreglo_id_raw:
            prefijo, pk_raw = arreglo_id_raw.split("-", 1)
            if prefijo == "F":
                tipo_item = "FLOR"
            elif prefijo == "P":
                tipo_item = "PRODUCTO"
            item_pk = pk_raw.strip()

        items.append(
            {
                "arreglo_id": arreglo_id_raw,
                "tipo_item": tipo_item,
                "item_pk": item_pk,
                "cantidad": cantidad_raw,
                "precio": precio_raw,
            }
        )

    return items


def _lock_item(tipo_item, item_pk):
    if tipo_item == "FLOR":
        return Flor.objects.select_for_update().get(pk=item_pk)
    return Producto.objects.select_for_update().get(pk=item_pk)


def _descontar_stock(tipo_item, item_pk, cantidad):
    item = _lock_item(tipo_item, item_pk)
    if item.cantidad < cantidad:
        raise ValueError(
            f"Stock insuficiente para {item.nombre}. "
            f"Disponible: {item.cantidad}, solicitado: {cantidad}."
        )
    item.cantidad -= cantidad
    item.save(update_fields=["cantidad"])
    crear_notificacion_stock(item.nombre, item.cantidad, "Venta")


def _devolver_stock(tipo_item, item_pk, cantidad):
    item = _lock_item(tipo_item, item_pk)
    item.cantidad += cantidad
    item.save(update_fields=["cantidad"])
    crear_notificacion_stock(item.nombre, item.cantidad, "Reversion de venta")


# ✅ Después — agrega LogoAE.png
def _get_logo_base64():
    candidatos = []
    for static_dir in getattr(settings, "STATICFILES_DIRS", []):
        candidatos.append(os.path.join(static_dir, "img", "LogoAE.png"))
        candidatos.append(os.path.join(static_dir, "img", "logo.png"))
        candidatos.append(os.path.join(static_dir, "img", "logo.jpg"))
    static_root = getattr(settings, "STATIC_ROOT", None)
    if static_root:
        candidatos.append(os.path.join(static_root, "img", "LogoAE.png"))
        candidatos.append(os.path.join(static_root, "img", "logo.png"))
        candidatos.append(os.path.join(static_root, "img", "logo.jpg"))
    for path in candidatos:
        if os.path.exists(path):
            with open(path, "rb") as f:
                return base64.b64encode(f.read()).decode("utf-8")
    return None


def _thin_border():
    side = Side(style="thin", color="D1D5DB")
    return Border(left=side, right=side, top=side, bottom=side)


def _anios_disponibles_ventas():
    anio_actual = timezone.now().year
    fechas      = Venta.objects.dates("fecha", "year", order="DESC")
    anios       = [d.year for d in fechas]
    if not anios:
        anios = [anio_actual]
    elif anio_actual not in anios:
        anios.insert(0, anio_actual)
    return anios, anio_actual


# ────────────────────────────────────────────────
#  LISTAR VENTAS
# ────────────────────────────────────────────────

@login_required
@panel_login_required
def listar_ventas(request):
    ventas = Venta.objects.select_related("cliente").prefetch_related(
        "detalles__flor", "detalles__producto"
    )

    q              = request.GET.get("q", "").strip()
    cliente_nombre = request.GET.get("cliente_nombre", "").strip()
    fecha_desde    = request.GET.get("fecha_desde", "").strip()
    precio_min     = request.GET.get("precio_min", "").strip()
    precio_max     = request.GET.get("precio_max", "").strip()

    def _parse_fecha(raw_fecha):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(raw_fecha, fmt).date()
            except ValueError:
                continue
        return None

    def _parse_decimal(raw_numero):
        valor = (raw_numero or "").strip()
        if not valor:
            return None
        try:
            if "," in valor:
                valor = valor.replace(".", "").replace(",", ".")
            return Decimal(valor)
        except (InvalidOperation, ValueError, TypeError):
            return None

    if q:
        filtros_q = (
            Q(cliente__nombre__icontains=q) |
            Q(cliente__apellido__icontains=q) |
            Q(cliente__documento__icontains=q)
        )
        if q.isdigit():
            filtros_q |= Q(id=int(q))
        ventas = ventas.filter(filtros_q)

    if cliente_nombre:
        partes = [p for p in cliente_nombre.split() if p]
        filtro = (
            Q(cliente__nombre__icontains=cliente_nombre) |
            Q(cliente__apellido__icontains=cliente_nombre)
        )
        if len(partes) >= 2:
            filtro |= (
                Q(cliente__nombre__icontains=partes[0]) &
                Q(cliente__apellido__icontains=" ".join(partes[1:]))
            )
        ventas = ventas.filter(filtro)

    if fecha_desde:
        fecha = _parse_fecha(fecha_desde)
        if fecha:
            ventas = ventas.filter(fecha__gte=fecha)
        else:
            fecha_desde = ""

    precio_min_val = _parse_decimal(precio_min)
    precio_max_val = _parse_decimal(precio_max)
    if precio_min and precio_min_val is None:
        precio_min = ""
    if precio_max and precio_max_val is None:
        precio_max = ""
    if precio_min_val is not None:
        ventas = ventas.filter(total__gte=precio_min_val)
    if precio_max_val is not None:
        ventas = ventas.filter(total__lte=precio_max_val)

    clientes       = Cliente.objects.all().order_by("nombre", "apellido")
    total_ventas   = ventas.count()
    monto_total    = ventas.aggregate(Sum("total"))["total__sum"] or 0
    total_clientes = ventas.values("cliente").distinct().count()

    hoy        = date.today()
    inicio_mes = hoy.replace(day=1)
    ventas_mes = ventas.filter(fecha__gte=inicio_mes, fecha__lte=hoy).count()

    anios_disponibles, anio_actual = _anios_disponibles_ventas()

    context = {
        "ventas":                ventas,
        "query":                 q,
        "clientes":              clientes,
        "cliente_nombre_filtro": cliente_nombre,
        "fecha_desde_filtro":    fecha_desde,
        "precio_min_filtro":     precio_min,
        "precio_max_filtro":     precio_max,
        "total_ventas":          total_ventas,
        "monto_total":           monto_total,
        "total_clientes":        total_clientes,
        "ventas_mes":            ventas_mes,
        "resultados_filtrados":  ventas.count(),
        "hay_filtros":           any([q, cliente_nombre, fecha_desde, precio_min, precio_max]),
        # ✅ Para el modal de reporte
        "anios_disponibles":     anios_disponibles,
        "anio_actual":           anio_actual,
        "tipos_venta":           Venta._meta.get_field("tipo_venta").choices,
    }
    return render(request, "ventas/listar_venta.html", context)


# ────────────────────────────────────────────────
#  REPORTE DE VENTAS
#  GET ?mes=&anio=&tipo_venta=&formato=html|excel
# ────────────────────────────────────────────────

@login_required
@panel_login_required
def reporte_venta(request):
    anio       = request.GET.get("anio", "").strip()
    mes        = request.GET.get("mes", "").strip()
    tipo_venta = request.GET.get("tipo_venta", "").strip()
    formato    = request.GET.get("formato", "html").strip()

    try:
        anio = int(anio)
    except (TypeError, ValueError):
        anio = timezone.now().year

    qs = (
        Venta.objects
        .filter(fecha__year=anio)
        .select_related("cliente")
        .prefetch_related("detalles")
    )

    # Filtro mes
    mes_nombre = None
    mes_int    = None
    if mes:
        try:
            mes_int    = int(mes)
            mes_nombre = MESES.get(mes_int, "")
            qs         = qs.filter(fecha__month=mes_int)
        except ValueError:
            mes = ""

    # Filtro tipo de venta
    tipo_nombre = TIPO_VENTA_DISPLAY.get(tipo_venta, "")
    if tipo_venta:
        qs = qs.filter(tipo_venta=tipo_venta)

    qs = qs.order_by("fecha")

    # Agregados
    # ✅ Después
    agg         = qs.aggregate(total=Sum("total"))
    monto_total = agg["total"] or Decimal("0")
    total_reg_calc = qs.count()
    promedio    = (monto_total / total_reg_calc) if total_reg_calc > 0 else Decimal("0")

    # Con / sin domicilio
    con_domicilio    = qs.filter(con_domicilio=True).count()
    sin_domicilio    = qs.filter(con_domicilio=False).count()
    total_envios     = qs.filter(con_domicilio=True).aggregate(
        t=Sum("precio_envio"))["t"] or Decimal("0")

    periodo   = f"{mes_nombre} {anio}".strip() if mes_nombre else str(anio)
    fecha_gen = datetime.now().strftime("%d/%m/%Y %H:%M")
    usuario   = (
        request.user.get_full_name() or request.user.username
        if request.user.is_authenticated else "—"
    )

    if formato == "excel":
        return _reporte_venta_excel(qs, periodo, monto_total)

    context = {
        "ventas":          qs,
        "periodo":         periodo,
        "mes_nombre":      mes_nombre,
        "mes":             mes or "",
        "anio":            anio,
        "tipo_venta":      tipo_venta,
        "tipo_nombre":     tipo_nombre,
        "monto_total":     monto_total,
        "promedio_venta":  promedio,
        "con_domicilio":   con_domicilio,
        "sin_domicilio":   sin_domicilio,
        "total_envios":    total_envios,
        "fecha_generacion":fecha_gen,
        "usuario":         usuario,
        "total_registros": total_reg_calc,
        "total_clientes":  qs.values("cliente").distinct().count(),
        "hay_filtros":     bool(mes_nombre or tipo_venta),
        "logo_base64":     _get_logo_base64(),
    }
    return render(request, "ventas/reporte_venta.html", context)


# ────────────────────────────────────────────────
#  EXCEL HELPER
# ────────────────────────────────────────────────

def _reporte_venta_excel(qs, periodo, total_general):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ventas"

    GREEN = "16A34A"
    LIGHT = "F0FDF4"

    ws.merge_cells("A1:H1")
    ws["A1"] = f"Reporte de Ventas — Artes & Estilos | Período: {periodo}"
    ws["A1"].font      = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill      = PatternFill("solid", fgColor=GREEN)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = ["#", "Fecha", "Cliente", "Tipo", "Forma Pago", "Ítems", "Domicilio", "Total"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font      = Font(bold=True, color="FFFFFF", size=10)
        cell.fill      = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = _thin_border()

    for i, v in enumerate(qs, start=3):
        fill = (PatternFill("solid", fgColor=LIGHT) if i % 2 == 0
                else PatternFill("solid", fgColor="FFFFFF"))
        cliente_nombre = (
            f"{v.cliente.nombre} {v.cliente.apellido}"
            if v.cliente else "N/A"
        )
        row_data = [
            v.id,
            v.fecha.strftime("%d/%m/%Y"),
            cliente_nombre,
            TIPO_VENTA_DISPLAY.get(v.tipo_venta, v.tipo_venta),
            v.get_forma_pago_display(),
            v.detalles.count(),
            "Sí" if v.con_domicilio else "No",
            float(v.total),
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.fill      = fill
            cell.border    = _thin_border()
            cell.alignment = Alignment(vertical="center")
            if col == 8:
                cell.number_format = '"$"#,##0.00'
                cell.alignment     = Alignment(horizontal="right", vertical="center")

    last = ws.max_row + 1
    ws.cell(row=last, column=6, value="TOTAL GENERAL").font = Font(bold=True)
    tot = ws.cell(row=last, column=8, value=float(total_general))
    tot.font          = Font(bold=True)
    tot.number_format = '"$"#,##0.00'
    tot.alignment     = Alignment(horizontal="right")
    tot.fill          = PatternFill("solid", fgColor="DCFCE7")
    for col in range(1, 9):
        ws.cell(row=last, column=col).border = _thin_border()

    # ✅ Letras directas — evita error MergedCell
    for col, w in zip(["A","B","C","D","E","F","G","H"], [6, 12, 28, 16, 16, 7, 10, 14]):
        ws.column_dimensions[col].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"reporte_ventas_{periodo.replace(' ', '_')}.xlsx"
    response = HttpResponse(
        buffer,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ────────────────────────────────────────────────
#  CRUD VENTAS (sin cambios)
# ────────────────────────────────────────────────

@login_required
@panel_login_required
def crear_venta(request):
    flores    = Flor.objects.all().order_by("nombre")
    productos = Producto.objects.all().order_by("nombre")
    mostrar_campos_domicilio = False

    if request.method == "POST":
        form = VentaForm(request.POST)
        mostrar_campos_domicilio = request.POST.get("con_domicilio") in ("on", "true", "1")
        try:
            detalles = _parse_detalles_venta(request)
        except ValueError as exc:
            messages.error(request, str(exc))
            return render(
                request,
                "ventas/agregar_venta.html",
                {
                    "form": form,
                    "flores": flores,
                    "productos": productos,
                    "mostrar_campos_domicilio": mostrar_campos_domicilio,
                    "posted_items": _obtener_items_posteados_venta(request),
                },
            )

        if form.is_valid():
            try:
                with transaction.atomic():
                    venta = form.save(commit=False)
                    venta.usuario = request.user
                    venta.total = Decimal("0")
                    venta.save()
                    for data in detalles:
                        detalle = DetalleVenta(
                            venta=venta, tipo_item=data["tipo_item"],
                            cantidad=data["cantidad"], precio=data["precio"],
                        )
                        if data["tipo_item"] == "FLOR":
                            detalle.flor = Flor.objects.get(pk=data["item_pk"])
                        else:
                            detalle.producto = Producto.objects.get(pk=data["item_pk"])
                        detalle.save()
                        _descontar_stock(data["tipo_item"], data["item_pk"], data["cantidad"])
                    venta.recalcular_totales()
                    venta.save(update_fields=["subtotal", "total"])
                    crear_notificacion(
                        categoria="movimiento", estilo="success",
                        titulo="Venta creada",
                        mensaje=f"Se registro la venta #{venta.id} con {len(detalles)} item(s).",
                    )
                messages.success(request, f"Venta #{venta.id} registrada correctamente.")
                return redirect("ventas:listar_venta")
            except (Flor.DoesNotExist, Producto.DoesNotExist):
                messages.error(request, "Uno de los items seleccionados ya no existe.")
            except ValueError as exc:
                messages.error(request, str(exc))
            except Exception as exc:
                messages.error(request, f"No se pudo registrar la venta: {exc}")
    else:
        form = VentaForm()

    return render(request, "ventas/agregar_venta.html", {
        "form": form, "flores": flores, "productos": productos,
        "mostrar_campos_domicilio": mostrar_campos_domicilio,
    })


@login_required
@panel_login_required
def editar_venta(request, pk):
    venta     = get_object_or_404(
        Venta.objects.prefetch_related("detalles__flor", "detalles__producto"), pk=pk
    )
    flores    = Flor.objects.all().order_by("nombre")
    productos = Producto.objects.all().order_by("nombre")
    mostrar_campos_domicilio = bool(venta.con_domicilio)

    if request.method == "POST":
        form = VentaForm(request.POST, instance=venta)
        mostrar_campos_domicilio = request.POST.get("con_domicilio") in ("on", "true", "1")
        try:
            nuevos_detalles = _parse_detalles_venta(request)
        except ValueError as exc:
            messages.error(request, str(exc))
            return render(
                request,
                "ventas/editar_venta.html",
                {
                    "form": form,
                    "venta": venta,
                    "detalles": venta.detalles.all(),
                    "flores": flores,
                    "productos": productos,
                    "mostrar_campos_domicilio": mostrar_campos_domicilio,
                    "posted_items": _obtener_items_posteados_venta(request),
                },
            )

        if form.is_valid():
            try:
                with transaction.atomic():
                    venta = form.save()
                    for detalle in list(venta.detalles.select_related("flor", "producto")):
                        item_pk = detalle.flor_id if detalle.tipo_item == "FLOR" else detalle.producto_id
                        if item_pk:
                            _devolver_stock(detalle.tipo_item, item_pk, detalle.cantidad)
                    venta.detalles.all().delete()
                    for data in nuevos_detalles:
                        detalle = DetalleVenta(
                            venta=venta, tipo_item=data["tipo_item"],
                            cantidad=data["cantidad"], precio=data["precio"],
                        )
                        if data["tipo_item"] == "FLOR":
                            detalle.flor = Flor.objects.get(pk=data["item_pk"])
                        else:
                            detalle.producto = Producto.objects.get(pk=data["item_pk"])
                        detalle.save()
                        _descontar_stock(data["tipo_item"], data["item_pk"], data["cantidad"])
                    venta.recalcular_totales()
                    venta.save(update_fields=["subtotal", "total"])
                    crear_notificacion(
                        categoria="movimiento", estilo="info",
                        titulo="Venta actualizada",
                        mensaje=f"Se actualizo la venta #{venta.id} con {len(nuevos_detalles)} item(s).",
                    )
                messages.success(request, f"Venta #{venta.id} actualizada correctamente.")
                return redirect("ventas:listar_venta")
            except (Flor.DoesNotExist, Producto.DoesNotExist):
                messages.error(request, "Uno de los items seleccionados ya no existe.")
            except ValueError as exc:
                messages.error(request, str(exc))
            except Exception as exc:
                messages.error(request, f"No se pudo actualizar la venta: {exc}")
    else:
        form = VentaForm(instance=venta)
    return render(
        request,
        "ventas/editar_venta.html",
        {
            "form": form,
            "venta": venta,
            "detalles": venta.detalles.all(),
            "flores": flores,
            "productos": productos,
            "mostrar_campos_domicilio": mostrar_campos_domicilio,
            "posted_items": _obtener_items_posteados_venta(request) if request.method == "POST" else [],
        },
    )


@login_required
@panel_login_required
def detalle_venta(request, pk):
    venta = get_object_or_404(
        Venta.objects.prefetch_related("detalles__flor", "detalles__producto").select_related("cliente", "usuario"),
        pk=pk,
    )
    venta.recalcular_totales()
    venta.save(update_fields=["subtotal", "total"])
    return render(request, "ventas/detalle_venta.html", {"venta": venta})


@login_required
@panel_login_required
def eliminar_venta(request, pk):
    venta = get_object_or_404(
        Venta.objects.prefetch_related("detalles__flor", "detalles__producto"), pk=pk
    )
    if request.method == "POST":
        try:
            with transaction.atomic():
                for detalle in list(venta.detalles.select_related("flor", "producto")):
                    item_pk = detalle.flor_id if detalle.tipo_item == "FLOR" else detalle.producto_id
                    if item_pk:
                        _devolver_stock(detalle.tipo_item, item_pk, detalle.cantidad)
                venta.delete()
                crear_notificacion(
                    categoria="movimiento", estilo="error",
                    titulo="Venta eliminada",
                    mensaje=f"Se elimino la venta #{pk}.",
                )
            messages.success(request, f"Venta #{pk} eliminada correctamente.")
            return redirect("ventas:listar_venta")
        except Exception as exc:
            messages.error(request, f"No se pudo eliminar la venta: {exc}")
    return render(request, "ventas/eliminar_venta.html", {"venta": venta})


@login_required
@panel_login_required
def buscar_cliente(request):
    q       = request.GET.get("q", "").strip()
    clientes = Cliente.objects.filter(nombre__icontains=q)[:10]
    return JsonResponse({"clientes": list(clientes.values("id", "nombre", "direccion"))})


@login_required
@panel_login_required
def buscar_arreglo(request):
    q         = request.GET.get("q", "").strip()
    flores    = Flor.objects.filter(nombre__icontains=q)[:10]
    productos = Producto.objects.filter(nombre__icontains=q)[:10]

    data = []
    for f in flores:
        data.append({
            "id": f"F-{f.id}", "nombre_flor": f.nombre,
            "tipo_producto": "Flor", "descripcion": f.descripcion,
            "precio": str(f.precio), "stock": f.cantidad,
            "imagen": f.imagen.url if f.imagen else "",
        })
    for p in productos:
        data.append({
            "id": f"P-{p.id}", "nombre_flor": p.nombre,
            "tipo_producto": "Producto", "descripcion": p.descripcion,
            "precio": str(p.precio), "stock": p.cantidad,
            "imagen": p.imagen.url if p.imagen else "",
        })
    return JsonResponse({"arreglos": data})